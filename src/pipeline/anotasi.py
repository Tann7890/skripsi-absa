import os
import random
from collections import Counter
from itertools import combinations
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

from pipeline.config import (
    DATA_ACUAN_FINAL,
    DATA_ACUAN_IDS,
    DAFTAR_DISKUSI,
    KALIMAT_BERSIH,
    KOLOM_DATA_ACUAN_FINAL,
    KOLOM_DISKUSI,
    KOLOM_KAPPA,
    KOLOM_TIDAK_SEPAKAT,
    PILIHAN_ASPEK,
    PILIHAN_SENTIMEN,
    PILIHAN_SET_KAPPA,
    PILIHAN_SET_LEMBAR,
    UJI_COBA_30,
    path_laporan_kappa,
    path_lembar_anotasi,
    path_tidak_sepakat,
)
from pipeline.data_acuan import cek_tidak_bocor
from pipeline.logbook import catat
from pipeline.parsing import sebaran
from pipeline.storage import baca_csv, tulis_csv

KOLOM_LEMBAR = ["No", "id_kalimat", "Kalimat", "label_aspek", "label_sentimen", "Catatan"]
LEBAR_KOLOM = {"A": 6, "B": 14, "C": 70, "D": 18, "E": 16, "F": 30}
PETUNJUK = [
    "Petunjuk singkat (baca Panduan Anotasi lengkap sebelum mulai)",
    "1. Kerjakan sendiri, jangan berdiskusi atau melihat lembar anotator lain.",
    "2. Untuk setiap kalimat pilih label_aspek lalu label_sentimen dari dropdown.",
    "3. label_aspek: tangibles, reliability, responsiveness, assurance, empathy, lainnya.",
    "4. label_sentimen: positif, negatif, netral; dinilai terhadap aspek yang dipilih.",
    "5. Kolom Catatan boleh diisi bila ragu; kolom lain tidak bisa dan tidak boleh diubah.",
    "6. Simpan tetap sebagai .xlsx lalu kirim ke peneliti.",
]
KOLOM_LEMBAR_ASPEK = ["No", "id_kalimat", "Kalimat", "label_aspek", "Catatan"]
LEBAR_KOLOM_ASPEK = {"A": 6, "B": 14, "C": 70, "D": 18, "E": 30}
PETUNJUK_LATIH = [
    "Petunjuk singkat (baca Panduan Anotasi lengkap sebelum mulai)",
    "1. Pilih label_aspek dari dropdown untuk setiap kalimat; data latih hanya diberi label aspek.",
    "2. label_aspek: tangibles, reliability, responsiveness, assurance, empathy, lainnya.",
    "3. Kolom Catatan boleh diisi bila ragu; kolom lain tidak bisa dan tidak boleh diubah.",
    "4. Simpan tetap sebagai .xlsx.",
]


def urutan_anotator(id_kalimat: list[str], nomor: int) -> list[str]:
    # acak tapi bisa diulang, tiap anotator dapat urutan beda
    urutan = sorted(id_kalimat)
    random.Random(42 + nomor).shuffle(urutan)
    return urutan


def buat_lembar(path: Path, baris: list[dict], sentimen: bool = True) -> None:
    header = KOLOM_LEMBAR if sentimen else KOLOM_LEMBAR_ASPEK
    lebar_kolom = LEBAR_KOLOM if sentimen else LEBAR_KOLOM_ASPEK
    teks_petunjuk = PETUNJUK if sentimen else PETUNJUK_LATIH
    pilihan_kolom = (("D", PILIHAN_ASPEK), ("E", PILIHAN_SENTIMEN)) if sentimen else (("D", PILIHAN_ASPEK),)
    n = len(baris)
    wb = Workbook()
    ws = wb.active
    ws.title = "Lembar"
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDDDDD")
    for nomor, r in enumerate(baris, 1):
        # tiga kolom awal terisi, sisanya kosong untuk diisi anotator
        ws.append([nomor, r["id_kalimat"], r["kalimat"]] + [""] * (len(header) - 3))
    for row in ws.iter_rows(min_row=1, max_row=n + 1):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for kolom, lebar in lebar_kolom.items():
        ws.column_dimensions[kolom].width = lebar
    ws.freeze_panes = "A2"
    # hanya kolom label yang boleh diubah, sisanya dikunci
    for row in ws.iter_rows(min_row=2, max_row=n + 1, min_col=4, max_col=len(header)):
        for c in row:
            c.protection = Protection(locked=False)
    # lembar dikunci tanpa password, cuma biar ga keedit
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    for kolom, pilihan in pilihan_kolom:
        validasi = DataValidation(
            type="list",
            formula1='"' + ",".join(pilihan) + '"',
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="Label tidak valid",
            error="Pilih label dari daftar",
        )
        validasi.add(f"{kolom}2:{kolom}{n + 1}")
        ws.add_data_validation(validasi)
    petunjuk = wb.create_sheet("Petunjuk")
    for i, teks in enumerate(teks_petunjuk, 1):
        petunjuk.cell(row=i, column=1, value=teks)
    petunjuk["A1"].font = Font(bold=True)
    petunjuk.column_dimensions["A"].width = 100
    path.parent.mkdir(parents=True, exist_ok=True)
    # simpan ke tmp lalu replace, lembar lama tidak rusak kalau gagal
    tmp = path.with_name(path.name + ".tmp.xlsx")
    wb.save(tmp)
    os.replace(tmp, path)


def jalankan_lembar_anotasi(cfg: dict, jenis: str) -> None:
    if jenis not in PILIHAN_SET_LEMBAR:
        raise ValueError(f"set harus salah satu dari {', '.join(PILIHAN_SET_LEMBAR)}")
    nomor_lembar = [0] if jenis == "peneliti" else list(range(1, cfg["jumlah_anotator"] + 1))
    paths = [path_lembar_anotasi(jenis, nomor) for nomor in nomor_lembar]
    for p in paths:
        if p.exists():
            raise FileExistsError(f"{p} sudah ada dan mungkin sudah diisi anotator; pindahkan dulu bila ingin membuat ulang")
    if not KALIMAT_BERSIH.exists():
        raise FileNotFoundError(f"{KALIMAT_BERSIH} belum ada; jalankan preprocessing dulu")
    # data acuan dan uji coba tidak boleh bocor ke lembar ini
    cek_tidak_bocor([])
    sumber = UJI_COBA_30 if jenis == "uji_coba" else DATA_ACUAN_IDS
    ids = [r["id_kalimat"] for r in baca_csv(sumber)]
    teks = {r["id_kalimat"]: r["kalimat"] for r in baca_csv(KALIMAT_BERSIH)}
    hilang = [i for i in ids if i not in teks]
    if hilang:
        raise ValueError(f"{len(hilang)} id_kalimat di {sumber} tidak ada di {KALIMAT_BERSIH}, misalnya {hilang[0]}")
    for nomor, path in zip(nomor_lembar, paths):
        buat_lembar(path, [{"id_kalimat": i, "kalimat": teks[i]} for i in urutan_anotator(ids, nomor)])
    if jenis == "peneliti":
        keputusan = (
            f"lembar peneliti dibuat ({len(ids)} kalimat data acuan) untuk dibandingkan dengan label konsensus (FR-45); "
            f"label peneliti tidak dipakai untuk kappa maupun label akhir; diisi bersamaan dengan anotator tanpa melihat lembar anotator"
        )
    else:
        keputusan = (
            f"{len(paths)} lembar {jenis} dibuat, masing-masing {len(ids)} kalimat dari {sumber.name}; "
            f"urutan acak berbeda per anotator (random.Random(42 + nomor)); kolom id_kalimat dan Kalimat dikunci; "
            f"tanpa merek, tipe, rating, dan prediksi model; peneliti bukan anotator"
        )
    catat("lembar_anotasi", keputusan, len(ids), len(paths), paths[0].parent)


def fleiss_kappa(penilaian: list[list[str]], kategori: tuple[str, ...]) -> float:
    m = len(penilaian[0]) if penilaian else 0
    if m < 2 or any(len(p) != m for p in penilaian):
        raise ValueError("fleiss_kappa butuh minimal 2 penilai dengan jumlah sama untuk setiap item")
    n_item = len(penilaian)
    hitungan = []
    for p in penilaian:
        hitung = Counter(p)
        hitungan.append([hitung[k] for k in kategori])
    # p_bar kesepakatan teramati, p_e kesepakatan karena kebetulan
    p_bar = sum((sum(n * n for n in baris) - m) / (m * (m - 1)) for baris in hitungan) / n_item
    p_e = sum((sum(baris[j] for baris in hitungan) / (n_item * m)) ** 2 for j in range(len(kategori)))
    # p_e 1 berarti semua label sama, kappa dianggap 1 bukan dibagi nol
    if p_e == 1:
        return 1.0
    return (p_bar - p_e) / (1 - p_e)


def cohen_kappa(a: list[str], b: list[str], kategori: tuple[str, ...]) -> float:
    n = len(a)
    p_o = sum(x == y for x, y in zip(a, b)) / n
    hitung_a = Counter(a)
    hitung_b = Counter(b)
    p_e = sum((hitung_a[k] / n) * (hitung_b[k] / n) for k in kategori)
    if p_e == 1:
        return 1.0
    return (p_o - p_e) / (1 - p_e)


def baca_lembar(path: Path, ids_harus: set[str], sentimen: bool = True) -> dict[str, dict]:
    if not path.exists():
        raise FileNotFoundError(f"{path} belum ada; kumpulkan lembar anotator dulu")
    # read_only cukup untuk membaca sel, lebih hemat memori
    wb = load_workbook(path, read_only=True)
    try:
        baris = list(wb["Lembar"].iter_rows(values_only=True))
    finally:
        wb.close()
    header = KOLOM_LEMBAR if sentimen else KOLOM_LEMBAR_ASPEK
    if not baris or list(baris[0][: len(header)]) != header:
        raise ValueError(f"{path.name}: header lembar tidak sesuai {header}")
    hasil: dict[str, dict] = {}
    masalah: list[str] = []
    for isi in baris[1:]:
        # baris kosong di akhir sheet dilewati
        if all(v is None or str(v).strip() == "" for v in isi[: len(header)]):
            continue
        no, id_mentah, _, aspek_mentah, *sisa = isi[: len(header)]
        id_kalimat = str(id_mentah or "").strip()
        aspek = str(aspek_mentah or "").strip().lower()
        if id_kalimat in hasil:
            masalah.append(f"No {no} id_kalimat ganda")
        if aspek not in PILIHAN_ASPEK:
            masalah.append(f"No {no} label_aspek kosong atau tidak valid")
        entri = {"aspek": aspek}
        if sentimen:
            label_sentimen = str(sisa[0] or "").strip().lower()
            if label_sentimen not in PILIHAN_SENTIMEN:
                masalah.append(f"No {no} label_sentimen kosong atau tidak valid")
            entri["sentimen"] = label_sentimen
        # catatan selalu kolom terakhir, baik set aspek maupun sentimen
        entri["catatan"] = str(sisa[-1] or "").strip()
        hasil[id_kalimat] = entri
    terbaca = set(hasil)
    masalah.extend(f"id_kalimat {i} hilang dari lembar" for i in sorted(ids_harus - terbaca))
    masalah.extend(f"id_kalimat {i} tidak ada di daftar sumber" for i in sorted(terbaca - ids_harus))
    if masalah:
        raise ValueError(f"{path.name}: {len(masalah)} masalah, misalnya: " + "; ".join(masalah[:5]))
    return hasil


def _baca_semua_lembar(cfg: dict, jenis: str) -> tuple[list[str], list[dict[str, dict]]]:
    # kappa dan label akhir baru dirancang untuk 3 anotator
    if cfg["jumlah_anotator"] != 3:
        raise ValueError("kappa dan label_akhir saat ini hanya mendukung 3 anotator")
    if jenis not in PILIHAN_SET_KAPPA:
        raise ValueError(f"set harus salah satu dari {', '.join(PILIHAN_SET_KAPPA)}")
    cek_tidak_bocor([])
    sumber = UJI_COBA_30 if jenis == "uji_coba" else DATA_ACUAN_IDS
    ids = [r["id_kalimat"] for r in baca_csv(sumber)]
    lembar = [baca_lembar(path_lembar_anotasi(jenis, nomor), set(ids)) for nomor in (1, 2, 3)]
    return ids, lembar


def hitung_kesepakatan(ids: list[str], lembar: list[dict[str, dict]]) -> list[dict]:
    hasil = []
    for label, kategori in (("aspek", PILIHAN_ASPEK), ("sentimen", PILIHAN_SENTIMEN)):
        penilaian = [[anotator[i][label] for anotator in lembar] for i in ids]
        hasil.append({"ukuran": "fleiss_kappa", "label": label, "anotator": "1,2,3", "nilai": round(fleiss_kappa(penilaian, kategori), 4)})
        # Cohen per pasangan, melihat anotator mana yang paling beda
        for a, b in combinations(range(3), 2):
            nilai = cohen_kappa([lembar[a][i][label] for i in ids], [lembar[b][i][label] for i in ids], kategori)
            hasil.append({"ukuran": "cohen_kappa", "label": label, "anotator": f"{a + 1},{b + 1}", "nilai": round(nilai, 4)})
        sepakat = sum(len({anotator[i][label] for anotator in lembar}) == 1 for i in ids) / len(ids)
        hasil.append({"ukuran": "persen_sepakat_semua", "label": label, "anotator": "1,2,3", "nilai": round(sepakat, 4)})
    return hasil


def daftar_tidak_sepakat(ids: list[str], lembar: list[dict[str, dict]], teks: dict[str, str]) -> list[dict]:
    daftar = []
    for i in ids:
        aspek = {anotator[i]["aspek"] for anotator in lembar}
        sentimen = {anotator[i]["sentimen"] for anotator in lembar}
        if len(aspek) == 1 and len(sentimen) == 1:
            continue
        baris = {"id_kalimat": i, "kalimat": teks[i]}
        for nomor, anotator in enumerate(lembar, 1):
            baris[f"aspek_{nomor}"] = anotator[i]["aspek"]
            baris[f"sentimen_{nomor}"] = anotator[i]["sentimen"]
            baris[f"catatan_{nomor}"] = anotator[i]["catatan"]
        daftar.append(baris)
    return daftar


def jalankan_kappa(cfg: dict, jenis: str) -> None:
    ids, lembar = _baca_semua_lembar(cfg, jenis)
    teks = {r["id_kalimat"]: r["kalimat"] for r in baca_csv(KALIMAT_BERSIH)}
    hasil = hitung_kesepakatan(ids, lembar)
    tulis_csv(path_laporan_kappa(jenis), hasil, KOLOM_KAPPA, timpa=True)
    tidak = daftar_tidak_sepakat(ids, lembar, teks)
    tulis_csv(path_tidak_sepakat(jenis), tidak, KOLOM_TIDAK_SEPAKAT, timpa=True)
    for r in hasil:
        print(f"{r['ukuran']}\t{r['label']}\t{r['anotator']}\t{r['nilai']}")
    fa, fs = [r["nilai"] for r in hasil if r["ukuran"] == "fleiss_kappa"]
    ambang = cfg["ambang_kappa"]
    if jenis == "uji_coba":
        status = "gambaran awal (30 kalimat, belum menentukan kelolosan)"
    elif fa >= ambang and fs >= ambang:
        status = "LOLOS"
    else:
        status = "TIDAK LOLOS: revisi panduan lalu kembali ke uji coba (Gambar 3.4)"
    catat(
        "kappa",
        f"set {jenis}: Fleiss Kappa aspek {fa:.4f}, sentimen {fs:.4f} (ambang {ambang}); {status}; "
        f"{len(tidak)} kalimat tidak sepakat penuh; Cohen per pasangan di {path_laporan_kappa(jenis).name}",
        len(ids),
        len(tidak),
        path_laporan_kappa(jenis),
    )


def tentukan_mayoritas(labels: list[str]) -> str | None:
    # None kalau tidak ada dua suara sama, masuk daftar diskusi
    label, n = Counter(labels).most_common(1)[0]
    return label if n >= 2 else None


def _baris_diskusi(perlu: list[str], teks: dict[str, str], lembar: list[dict[str, dict]], aspek: dict[str, str | None], sentimen: dict[str, str | None]) -> list[dict]:
    daftar = []
    for i in perlu:
        baris = {"id_kalimat": i, "kalimat": teks[i]}
        for nomor, anotator in enumerate(lembar, 1):
            baris[f"aspek_{nomor}"] = anotator[i]["aspek"]
            baris[f"sentimen_{nomor}"] = anotator[i]["sentimen"]
        baris["perlu_diskusi"] = ",".join(k for k, v in (("aspek", aspek[i]), ("sentimen", sentimen[i])) if v is None)
        daftar.append(baris)
    return daftar


def _isi_dari_diskusi(perlu: list[str], aspek: dict[str, str | None], sentimen: dict[str, str | None]) -> list[str]:
    diskusi = {r["id_kalimat"]: r for r in baca_csv(DAFTAR_DISKUSI)}
    masalah = []
    for i in perlu:
        if i not in diskusi:
            masalah.append(f"id_kalimat {i} tidak ada di {DAFTAR_DISKUSI.name}")
            continue
        for kunci, hasil, pilihan in (("aspek", aspek, PILIHAN_ASPEK), ("sentimen", sentimen, PILIHAN_SENTIMEN)):
            if hasil[i] is not None:
                continue
            nilai = diskusi[i][f"label_{kunci}_diskusi"].strip().lower()
            if nilai in pilihan:
                hasil[i] = nilai
            else:
                masalah.append(f"id_kalimat {i} label_{kunci}_diskusi kosong atau tidak valid")
    return masalah


def jalankan_label_akhir(cfg: dict) -> None:
    if DATA_ACUAN_FINAL.exists():
        raise FileExistsError(f"{DATA_ACUAN_FINAL} sudah ada dan tidak boleh ditimpa")
    ids, lembar = _baca_semua_lembar(cfg, "acuan")
    fa = fleiss_kappa([[anotator[i]["aspek"] for anotator in lembar] for i in ids], PILIHAN_ASPEK)
    fs = fleiss_kappa([[anotator[i]["sentimen"] for anotator in lembar] for i in ids], PILIHAN_SENTIMEN)
    ambang = cfg["ambang_kappa"]
    if fa < ambang or fs < ambang:
        raise ValueError(f"Fleiss Kappa aspek {fa:.4f}, sentimen {fs:.4f} di bawah ambang {ambang}; revisi panduan lalu kembali ke uji coba")
    aspek = {i: tentukan_mayoritas([anotator[i]["aspek"] for anotator in lembar]) for i in ids}
    sentimen = {i: tentukan_mayoritas([anotator[i]["sentimen"] for anotator in lembar]) for i in ids}
    perlu = [i for i in ids if aspek[i] is None or sentimen[i] is None]
    # daftar diskusi hanya diekspor sekali, jangan timpa isian peneliti
    if perlu and not DAFTAR_DISKUSI.exists():
        teks = {r["id_kalimat"]: r["kalimat"] for r in baca_csv(KALIMAT_BERSIH)}
        tulis_csv(DAFTAR_DISKUSI, _baris_diskusi(perlu, teks, lembar, aspek, sentimen), KOLOM_DISKUSI)
        catat(
            "label_akhir",
            f"{len(perlu)} kalimat tanpa suara terbanyak diekspor ke {DAFTAR_DISKUSI.name}; isi label hasil diskusi lalu jalankan label_akhir lagi",
            len(ids),
            len(perlu),
            DAFTAR_DISKUSI,
        )
        return
    # salinan supaya kolom sumber tetap bisa membedakan mayoritas dan diskusi
    final_aspek = dict(aspek)
    final_sentimen = dict(sentimen)
    if perlu:
        masalah = _isi_dari_diskusi(perlu, final_aspek, final_sentimen)
        if masalah:
            raise ValueError(f"{DAFTAR_DISKUSI.name}: {len(masalah)} masalah, misalnya: " + "; ".join(masalah[:5]))
    rows = [
        {
            "id_kalimat": i,
            "label_aspek": final_aspek[i],
            "label_sentimen": final_sentimen[i],
            "sumber_aspek": "mayoritas" if aspek[i] is not None else "diskusi",
            "sumber_sentimen": "mayoritas" if sentimen[i] is not None else "diskusi",
        }
        for i in ids
    ]
    tulis_csv(DATA_ACUAN_FINAL, rows, KOLOM_DATA_ACUAN_FINAL)
    a_m = sum(r["sumber_aspek"] == "mayoritas" for r in rows)
    s_m = sum(r["sumber_sentimen"] == "mayoritas" for r in rows)
    catat(
        "label_akhir",
        f"{len(rows)} label akhir data acuan; aspek: mayoritas {a_m}, diskusi {len(rows) - a_m}; "
        f"sentimen: mayoritas {s_m}, diskusi {len(rows) - s_m}; "
        f"sebaran aspek: {sebaran(rows, 'label_aspek')}; sebaran sentimen: {sebaran(rows, 'label_sentimen')}",
        len(ids),
        len(rows),
        DATA_ACUAN_FINAL,
    )
